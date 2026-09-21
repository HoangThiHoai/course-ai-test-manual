# -*- coding: utf-8 -*-
"""Dựng file Excel test case MỚI theo đúng format 1 sheet trong file TC mẫu.

Không sửa file mẫu / Google Sheet: script mở bản tải về, lấy 1 sheet làm khuôn
(giữ nguyên khối header, style, độ rộng cột), xoá nội dung cũ, ghi TC mới,
bỏ các sheet khác rồi lưu ra file mới.

Cách dùng:
    python build_tc_excel.py --template tcs.xlsx --template-sheet "QLĐH_Hủy đơn hàng" \
        --data tcdata.py --out TCs_<module>_<muc>.xlsx \
        --sheet-name "QLĐH_Chi tiết hoàn tiền" \
        --screen "Quản lý đơn hàng_Chi tiết hoàn tiền (SRS mục 1.3.4)" \
        --doc-link "SRS ... mục 1.3.4 | Figma ... node ..."

File --data là module Python khai báo:
    ROWS = [
      ('S', 'Tên section'),                       # dòng nhóm màn hình (nền màu)
      ('P', 'Pre-condition: \\nBước 1: ...'),      # dòng điều kiện tiên quyết
      ('T', chức_năng|None, mục_đích, các_bước, kết_quả, dữ_liệu, ghi_chú),
    ]
    TECHNIQUES = [   # tuỳ chọn -> sinh sheet "Kỹ thuật thiết kế TC"
      {'title': '1. BẢNG QUYẾT ĐỊNH ...', 'header': [...], 'rows': [[...], ...]},
    ]
chức_năng = None nghĩa là cùng nhóm với TC phía trên (ô cột Chức năng được merge).
"""
import argparse
import copy
import importlib.util
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


def load_data(path):
    spec = importlib.util.spec_from_file_location("tcdata", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.ROWS, getattr(mod, "TECHNIQUES", [])


def validate(rows):
    errors, seen, need_func = [], {}, True
    for i, item in enumerate(rows):
        kind = item[0]
        if kind in ("S", "P"):
            need_func = kind == "S" or need_func
            continue
        if kind != "T" or len(item) != 7:
            errors.append(f"Phần tử #{i}: sai cấu trúc, cần ('T', func, purpose, steps, expected, data, note)")
            continue
        _, func, purpose, steps, expected, _, _ = item
        if need_func and not func:
            errors.append(f"Phần tử #{i} ({purpose}): TC đầu tiên sau section phải có tên Chức năng")
        need_func = False
        for label, val in (("Mục đích", purpose), ("Các bước", steps), ("Kết quả mong muốn", expected)):
            if not str(val or "").strip():
                errors.append(f"Phần tử #{i}: thiếu {label}")
        key = (purpose, steps)
        if key in seen:
            errors.append(f"Phần tử #{i}: trùng hoàn toàn với phần tử #{seen[key]} ({purpose})")
        seen[key] = i
    return errors


def style_of(cell):
    return tuple(copy.copy(x) for x in (cell.font, cell.border, cell.alignment, cell.fill))


def apply(cell, sty):
    cell.font, cell.border, cell.alignment, cell.fill = (copy.copy(s) for s in sty)


def find_col(ws, header_row, *names):
    for c in ws[header_row]:
        if c.value and str(c.value).strip() in names:
            return c.column
    return None


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--template-sheet", required=True)
    ap.add_argument("--data", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--sheet-name", required=True)
    ap.add_argument("--screen", required=True, help="Giá trị ô 'Tên màn hình/Tên chức năng'")
    ap.add_argument("--doc-link", default="", help="Giá trị ô 'Link tài liệu'")
    ap.add_argument("--prefix", default="TC-", help="Tiền tố ID, mặc định TC-")
    args = ap.parse_args()

    rows, techniques = load_data(args.data)
    errs = validate(rows)
    if errs:
        print("[LỖI] Dữ liệu TC chưa hợp lệ:")
        for e in errs:
            print("  -", e)
        sys.exit(1)

    wb = openpyxl.load_workbook(args.template)
    ws = wb[args.template_sheet]
    header_row = next((c.row for c in ws["A"] if str(c.value).strip() == "ID"), None)
    if not header_row:
        sys.exit("[LỖI] Không tìm thấy dòng tiêu đề cột (ô cột A = 'ID') trong sheet mẫu")
    last_col = max(c.column for c in ws[header_row] if c.value)

    # Dòng tiêu đề có thể merge 2 dòng -> nội dung bắt đầu sau vùng merge
    content_start = header_row + 1
    for m in ws.merged_cells.ranges:
        if m.min_row == header_row and m.min_col == 1:
            content_start = m.max_row + 1

    sec_row = pre_row = tc_row = None
    for r in range(content_start, ws.max_row + 1):
        v = str(ws.cell(row=r, column=1).value or "").strip()
        if not v:
            continue
        if v.lower().startswith("pre-condition") and pre_row is None:
            pre_row = r
        elif (re.match(r"^[A-Za-z_]*TC[-_ ]?\d+", v) or v.lower().startswith("=")) and tc_row is None:
            tc_row = r
        elif sec_row is None and tc_row is None and not v.lower().startswith("pre-condition"):
            sec_row = r
        if sec_row and pre_row and tc_row:
            break
    if not tc_row:
        sys.exit("[LỖI] Sheet mẫu không có dòng TC nào để lấy style")
    sec_row = sec_row or tc_row
    pre_row = pre_row or tc_row

    def merge_width(row):
        for m in ws.merged_cells.ranges:
            if m.min_row == row and m.min_col == 1 and m.max_row == row:
                return m.max_col
        return 1

    sec_span, pre_span = merge_width(sec_row), merge_width(pre_row)
    # File mẫu đánh ID bằng công thức (=$E$4&"-"&TEXT(COUNTA(...),"00")) thì giữ nguyên cách đó
    id_is_formula = str(ws.cell(row=tc_row, column=1).value or "").startswith("=")
    cols = range(1, last_col + 1)
    sty_tc = {c: style_of(ws.cell(row=tc_row, column=c)) for c in cols}
    sty_sec = {c: style_of(ws.cell(row=sec_row, column=c)) for c in cols}
    sty_pre = {c: style_of(ws.cell(row=pre_row, column=c)) for c in cols}

    col_func = find_col(ws, header_row, "Chức năng") or 2
    col_purpose = find_col(ws, header_row, "Mục đích") or 3
    col_steps = find_col(ws, header_row, "Các bước thực hiện kiểm thử") or 4
    col_expected = find_col(ws, header_row, "Kết quả mong muốn") or 5
    col_data = find_col(ws, header_row, "Dữ liệu kiểm thử") or 6
    col_status = find_col(ws, header_row, "Trạng thái")
    col_note = find_col(ws, header_row, "Ghi chú")

    for m in list(ws.merged_cells.ranges):
        if m.min_row >= content_start:
            ws.unmerge_cells(str(m))
    ws.data_validations.dataValidation = []
    ws.delete_rows(content_start, ws.max_row - content_start + 1)

    # Vị trí các ô trong khối header, tìm theo nhãn ở cột bên trái
    header_cells = {}
    label_keys = {"tên màn hình": "screen", "mã testcase": "code", "số testcase đạt": "pass",
                  "không đạt": "fail", "chưa test": "untested", "tổng số testcase": "total",
                  "link tài liệu": "link"}
    for row in ws.iter_rows(min_row=1, max_row=header_row - 1):
        for c in row:
            txt = str(c.value or "").lower()
            for key, name in label_keys.items():
                if key in txt and name not in header_cells:
                    header_cells[name] = ws.cell(row=c.row, column=c.column + 1)

    L = openpyxl.utils.get_column_letter
    exp_letter = L(col_expected)
    code_cell = header_cells.get("code")
    code_ref = f"${code_cell.column_letter}${code_cell.row}" if code_cell else None

    r, tc_no, group_start, first_tc = content_start, 0, None, None

    def close_group(end_row):
        nonlocal group_start
        if group_start and end_row > group_start:
            ws.merge_cells(start_row=group_start, start_column=col_func, end_row=end_row, end_column=col_func)
        group_start = None

    for item in rows:
        if item[0] == "S":
            close_group(r - 1)
            for c in cols:
                apply(ws.cell(row=r, column=c), sty_sec[c])
            ws.cell(row=r, column=1, value=item[1])
            if sec_span > 1:
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=sec_span)
            ws.row_dimensions[r].height = 18
        elif item[0] == "P":
            for c in cols:
                apply(ws.cell(row=r, column=c), sty_pre[c])
            ws.cell(row=r, column=1, value=item[1])
            if pre_span > 1:
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=pre_span)
            ws.row_dimensions[r].height = max(45, 16 * (item[1].count("\n") + 1))
        else:
            _, func, purpose, steps, expected, data, note = item
            for c in cols:
                apply(ws.cell(row=r, column=c), sty_tc[c])
            tc_no += 1
            first_tc = first_tc or r
            if id_is_formula and code_ref:
                ws.cell(row=r, column=1,
                        value=f'={code_ref}&"-"&TEXT(COUNTA(${exp_letter}${first_tc}:{exp_letter}{r}),"00")')
            else:
                ws.cell(row=r, column=1, value=f"{args.prefix}{tc_no:02d}")
            if func:
                close_group(r - 1)
                ws.cell(row=r, column=col_func, value=func)
                group_start = r
            ws.cell(row=r, column=col_purpose, value=purpose)
            ws.cell(row=r, column=col_steps, value=steps)
            ws.cell(row=r, column=col_expected, value=expected)
            if data:
                ws.cell(row=r, column=col_data, value=data)
            if note and col_note:
                ws.cell(row=r, column=col_note, value=note)
            lines = max(steps.count("\n"), expected.count("\n")) + 1
            ws.row_dimensions[r].height = max(45, min(15 * (lines + 1), 240))
        r += 1
    close_group(r - 1)
    last_row = r - 1

    if col_status and first_tc:
        dv = DataValidation(type="list", formula1='"Pass,Fail,N/A"', allow_blank=True)
        ws.add_data_validation(dv)
        letter = openpyxl.utils.get_column_letter(col_status)
        dv.add(f"{letter}{first_tc}:{letter}{last_row}")

    # Khối header: thống kê bằng công thức để tự cập nhật khi tester chấm Trạng thái
    if "screen" in header_cells:
        header_cells["screen"].value = args.screen
    if "link" in header_cells and args.doc_link:
        header_cells["link"].value = args.doc_link
    if first_tc and col_status:
        st = L(col_status)
        st_rng = f"{st}{first_tc}:{st}{last_row}"
        exp_rng = f"{exp_letter}{first_tc}:{exp_letter}{last_row}"
        ref = {k: c.coordinate for k, c in header_cells.items()}
        formulas = {
            "pass": f'=COUNTIF({st_rng},"Pass")',
            "fail": f'=COUNTIF({st_rng},"Fail")',
            "total": f"=COUNTA({exp_rng})",
        }
        if {"total", "pass", "fail"} <= ref.keys():
            formulas["untested"] = f"={ref['total']}-{ref['pass']}-{ref['fail']}"
        for name, formula in formulas.items():
            if name in header_cells:
                header_cells[name].value = formula

    ws.title = args.sheet_name
    for name in list(wb.sheetnames):
        if name != ws.title:
            del wb[name]

    if techniques:
        build_technique_sheet(wb, techniques)

    wb.save(args.out)
    print(f"[OK] {args.out}: {tc_no} TC (dòng {content_start}..{last_row})"
          + (f", kèm sheet 'Kỹ thuật thiết kế TC' ({len(techniques)} bảng)" if techniques else ""))


def build_technique_sheet(wb, techniques):
    ws = wb.create_sheet("Kỹ thuật thiết kế TC")
    thin = Side(style="thin")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    f_title = Font(name="Times New Roman", size=13, bold=True)
    f_hdr = Font(name="Times New Roman", size=12, bold=True)
    f_txt = Font(name="Times New Roman", size=12)
    fill = PatternFill("solid", fgColor="FFD8D8D8")
    width = max(len(t["header"]) for t in techniques)
    for i in range(1, width + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 34 if i == 1 else 22

    r = 1
    for block in techniques:
        ws.cell(row=r, column=1, value=block["title"]).font = f_title
        r += 2
        for line, is_hdr in [(block["header"], True)] + [(x, False) for x in block["rows"]]:
            for i, v in enumerate(line, start=1):
                c = ws.cell(row=r, column=i, value=v)
                c.border = bd
                c.font = f_hdr if is_hdr else f_txt
                c.alignment = Alignment(horizontal="center" if is_hdr else "left",
                                        vertical="center", wrap_text=True)
                if is_hdr:
                    c.fill = fill
            ws.row_dimensions[r].height = 45 if is_hdr else 32
            r += 1
        r += 2


if __name__ == "__main__":
    main()
