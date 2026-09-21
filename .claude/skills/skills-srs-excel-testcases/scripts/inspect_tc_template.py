# -*- coding: utf-8 -*-
"""Đọc file test case mẫu (.xlsx) để học format trước khi viết TC mới.

Cách dùng:
    python inspect_tc_template.py <tcs.xlsx>                       # liệt kê sheet
    python inspect_tc_template.py <tcs.xlsx> --sheet "<tên sheet>" # soi 1 sheet
    python inspect_tc_template.py <tcs.xlsx> --sheet "<...>" --rows 80

In ra: khối header, dòng tiêu đề cột, dòng section / Pre-condition / TC mẫu,
merge, độ rộng cột, data validation, style (font/màu nền) -> đủ để tái tạo format.
"""
import argparse

import openpyxl


def short(v, n=300):
    s = str(v).replace("\n", "\\n")
    return s if len(s) <= n else s[:n] + "…"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sheet")
    ap.add_argument("--rows", type=int, default=40, help="Số dòng nội dung in ra sau dòng tiêu đề")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    if not args.sheet:
        print("Danh sách sheet (tên | số dòng | số cột):")
        for ws in wb.worksheets:
            print(f"  {ws.title} | {ws.max_row} | {ws.max_column}")
        print("\nChọn 1-2 sheet cùng dạng màn hình với mục cần viết rồi chạy lại với --sheet.")
        return

    ws = wb[args.sheet]
    header_row = next((c.row for c in ws["A"] if str(c.value).strip() == "ID"), None)
    print(f"== Sheet: {ws.title} | dòng tiêu đề cột: {header_row}")

    print("\n-- Khối header phía trên --")
    for row in ws.iter_rows(min_row=1, max_row=(header_row or 12) - 1):
        vals = [f"{c.coordinate}={short(c.value, 120)}" for c in row if c.value not in (None, "")]
        if vals:
            print("  " + " || ".join(vals))

    if header_row:
        print("\n-- Tiêu đề cột --")
        print("  " + " | ".join(f"{c.column_letter}:{c.value}" for c in ws[header_row] if c.value))
        print("\n-- Độ rộng cột --")
        print("  " + " ".join(f"{k}={v.width}" for k, v in sorted(ws.column_dimensions.items()) if v.width))

        print(f"\n-- Nội dung {args.rows} dòng đầu --")
        for row in ws.iter_rows(min_row=header_row + 1, max_row=header_row + args.rows):
            vals = [f"{c.coordinate}={short(c.value)}" for c in row if c.value not in (None, "")]
            if vals:
                print("  " + " || ".join(vals))

    print("\n-- Merge (30 vùng đầu) --")
    print("  " + ", ".join(str(m) for m in list(ws.merged_cells.ranges)[:30]))
    print("\n-- Data validation --")
    for dv in ws.data_validations.dataValidation:
        print(f"  {dv.sqref} type={dv.type} formula={dv.formula1}")

    if header_row:
        print("\n-- Style các ô mốc --")
        wb2 = openpyxl.load_workbook(args.xlsx)
        ws2 = wb2[args.sheet]
        for r in range(header_row, min(header_row + 6, ws2.max_row) + 1):
            c = ws2.cell(row=r, column=1)
            print(f"  A{r}: font={c.font.name} {c.font.sz} bold={c.font.b} fill={c.fill.fgColor.rgb} "
                  f"wrap={c.alignment.wrap_text} border={c.border.left.style} | {short(c.value, 50)}")


if __name__ == "__main__":
    main()
