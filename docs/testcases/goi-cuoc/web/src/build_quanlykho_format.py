# -*- coding: utf-8 -*-
"""Dựng file Excel TC theo khuôn sheet "Quanlykho" (file TC dự án) — header nhãn cột A / giá trị cột C,
2 cấp section (UC nền xanh merge A:C, nhóm con merge A:K), ID bằng công thức đếm cột D.

Cách dùng:
    python build_quanlykho_format.py --template tcs.xlsx --data tcdata_qlgc_uc1_uc2.py --out <file.xlsx> \
        --sheet-name "Quanlygoicuoc" --screen "..." --doc-link "..." --created "24/09/2026"
"""
import argparse
import copy
import importlib.util
import os
import sys

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

SKILL_SCRIPTS = os.path.join(os.path.dirname(__file__), '..', '..', '..', '..', '..',
                             '.claude', 'skills', 'skills-srs-excel-testcases', 'scripts')
sys.path.insert(0, os.path.abspath(SKILL_SCRIPTS))
from build_tc_excel import build_technique_sheet  # noqa: E402

TEMPLATE_SHEET = 'Quanlykho'
ROW_UC, ROW_PRE, ROW_TC, ROW_GROUP = 12, 13, 14, 21
CONTENT_START = 12
LAST_COL = 13  # A..M
COL_FUNC, COL_PURPOSE, COL_STEPS, COL_EXP, COL_DATA, COL_CREATED, COL_STATUS, COL_NOTE = 2, 3, 4, 5, 6, 7, 11, 13


def load(path):
    spec = importlib.util.spec_from_file_location('tcdata', path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.ROWS, getattr(mod, 'TECHNIQUES', [])


def validate(rows):
    errors, seen, need_func, section, cur_func = [], {}, True, '', ''
    for i, item in enumerate(rows):
        if item[0] in ('S', 'G', 'P'):
            need_func = need_func or item[0] in ('S', 'G')
            if item[0] == 'S':
                section = item[1]
            continue
        if item[0] != 'T' or len(item) != 7:
            errors.append(f'#{i}: sai cấu trúc')
            continue
        _, func, purpose, steps, expected, _, _ = item
        if need_func and not func:
            errors.append(f'#{i} ({purpose}): TC đầu nhóm thiếu Chức năng')
        need_func = False
        if not (purpose and steps and expected):
            errors.append(f'#{i}: thiếu Mục đích/Bước/Kết quả')
        cur_func = func or cur_func
        key = (section, cur_func, purpose, steps)
        if key in seen:
            errors.append(f'#{i}: trùng với #{seen[key]} ({purpose})')
        seen[key] = i
    return errors


def style_of(cell):
    return tuple(copy.copy(x) for x in (cell.font, cell.border, cell.alignment, cell.fill, cell.number_format))


def apply(cell, sty):
    cell.font, cell.border, cell.alignment, cell.fill = (copy.copy(s) for s in sty[:4])
    cell.number_format = sty[4]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--template', required=True)
    ap.add_argument('--data', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--sheet-name', required=True)
    ap.add_argument('--screen', required=True)
    ap.add_argument('--doc-link', default='')
    ap.add_argument('--created', default='')
    args = ap.parse_args()

    rows, techniques = load(args.data)
    errs = validate(rows)
    if errs:
        sys.exit('[LỖI]\n' + '\n'.join(errs))

    wb = openpyxl.load_workbook(args.template)
    ws = wb[TEMPLATE_SHEET]
    cols = range(1, LAST_COL + 1)
    sty = {k: {c: style_of(ws.cell(row=r, column=c)) for c in cols}
           for k, r in (('S', ROW_UC), ('P', ROW_PRE), ('T', ROW_TC), ('G', ROW_GROUP))}
    span = {'S': 3, 'G': 11, 'P': 13}

    for m in list(ws.merged_cells.ranges):
        if m.min_row >= CONTENT_START:
            ws.unmerge_cells(str(m))
    ws.data_validations.dataValidation = []
    ws.delete_rows(CONTENT_START, ws.max_row - CONTENT_START + 1)

    r, group_start, first_tc, n = CONTENT_START, None, None, 0

    def close_group(end):
        nonlocal group_start
        if group_start and end > group_start:
            ws.merge_cells(start_row=group_start, start_column=COL_FUNC, end_row=end, end_column=COL_FUNC)
        group_start = None

    for item in rows:
        kind = item[0]
        for c in cols:
            apply(ws.cell(row=r, column=c), sty[kind][c])
        if kind in span:
            close_group(r - 1)
            ws.cell(row=r, column=1, value=item[1])
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span[kind])
            ws.row_dimensions[r].height = max(20, 16 * (item[1].count('\n') + 1)) if kind == 'P' else 20
        else:
            _, func, purpose, steps, expected, data, note = item
            n += 1
            first_tc = first_tc or r
            ws.cell(row=r, column=1, value=f'=$C$4&"-"&TEXT(COUNTA($D${CONTENT_START}:D{r}),"00")')
            if func:
                close_group(r - 1)
                ws.cell(row=r, column=COL_FUNC, value=func)
                group_start = r
            ws.cell(row=r, column=COL_PURPOSE, value=purpose)
            ws.cell(row=r, column=COL_STEPS, value=steps)
            ws.cell(row=r, column=COL_EXP, value=expected)
            if data:
                ws.cell(row=r, column=COL_DATA, value=data)
            if args.created:
                ws.cell(row=r, column=COL_CREATED, value=args.created)
            if note:
                ws.cell(row=r, column=COL_NOTE, value=note)
            lines = max(steps.count('\n'), expected.count('\n'), len(expected) // 45, len(note or '') // 30) + 1
            ws.row_dimensions[r].height = max(32, min(16 * lines, 260))
        r += 1
    close_group(r - 1)
    last = r - 1

    dv = DataValidation(type='list', formula1='"Pass,Fail,Pending,N/A"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'K{first_tc}:K{last}')

    ws['C2'] = args.screen
    ws['C3'] = args.doc_link
    ws['C6'] = f'=COUNTIF(A{CONTENT_START}:A{last},"*TC*")'
    ws['C7'] = f'=COUNTIF($K${first_tc}:$K${last},"Pass")'
    ws['C8'] = f'=COUNTIF($K${first_tc}:$K${last},"Fail")'
    ws['C9'] = f'=COUNTIF($K${first_tc}:$K${last},"N/A")'

    ws.title = args.sheet_name
    for name in list(wb.sheetnames):
        if name != ws.title:
            del wb[name]
    if techniques:
        build_technique_sheet(wb, techniques)
    wb.save(args.out)
    print(f'[OK] {args.out}: {n} TC (dòng {CONTENT_START}..{last})')


if __name__ == '__main__':
    main()
