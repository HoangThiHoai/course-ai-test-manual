# -*- coding: utf-8 -*-
"""Dựng bản NHÁP: chép nguyên sheet TC của dự án rồi CHÈN các TC bổ sung vào giữa, đúng vị trí sẽ dán vào
Google Sheet. Không sửa file gốc / Google Sheet.

Giữ đúng format sheet "CMS-Quản lý gói cước" (cột A..G: ID · Chức năng · Mục đích · Các bước · Kết quả · Trạng thái
· Ghi chú): ID bằng công thức đếm cột E, nhóm merge A:F, Pre-condition merge A:G, cột Chức năng merge dọc,
dropdown Trạng thái. Dòng bổ sung tô nền vàng nhạt để review (xoá màu trước khi dán nếu không cần).

Cách dùng:
    python build_insert_bosung.py --template tcs.xlsx --sheet "CMS-Quản lý gói cước" \
        --before "Kiểm tra ngoại lệ & phi chức năng" --data tcdata_qlgc_caidatbang_boloc.py \
        --out TCs_....xlsx --date 02/10/2026
"""
import argparse
import copy
import importlib.util
import os
import re
import sys

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.styles import PatternFill
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.datavalidation import DataValidation

SKILL_SCRIPTS = os.path.join(os.path.dirname(__file__), '..', '..', '..', '..', '..',
                             '.claude', 'skills', 'skills-srs-excel-testcases', 'scripts')
sys.path.insert(0, os.path.abspath(SKILL_SCRIPTS))
from build_tc_excel import build_technique_sheet  # noqa: E402

COL_ID, COL_FUNC, COL_PURPOSE, COL_STEPS, COL_EXP, COL_STATUS, COL_NOTE = 1, 2, 3, 4, 5, 6, 7
LAST_COL = 7
FIRST_COUNT_ROW = 15  # công thức ID của file mẫu đếm từ $E$15
NEW_FILL = PatternFill('solid', fgColor='FFFFF2CC')


def load(path):
    spec = importlib.util.spec_from_file_location('tcdata', path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.ROWS, getattr(mod, 'TECHNIQUES', [])


def validate(rows):
    errors, seen, need_func, section = [], {}, True, ''
    for i, item in enumerate(rows):
        if item[0] in ('S', 'P'):
            need_func = need_func or item[0] == 'S'
            section = item[1] if item[0] == 'S' else section
            continue
        if item[0] != 'T' or len(item) not in (6, 7):
            errors.append(f'#{i}: sai cấu trúc')
            continue
        _, func, purpose, steps, expected = item[:5]
        if need_func and not func:
            errors.append(f'#{i} ({purpose}): TC đầu nhóm thiếu Chức năng')
        need_func = False
        if not (purpose and steps and expected):
            errors.append(f'#{i}: thiếu Mục đích/Bước/Kết quả')
        key = (section, purpose, steps)
        if key in seen:
            errors.append(f'#{i}: trùng với #{seen[key]} ({purpose})')
        seen[key] = i
    return errors


def indent(text):
    """File mẫu viết xuống dòng dạng '\\n ' (có 1 dấu cách đầu dòng)."""
    return re.sub(r'\n(?! )', '\n ', text)


def style_of(cell):
    return tuple(copy.copy(x) for x in (cell.font, cell.border, cell.alignment, cell.fill, cell.number_format))


def apply(cell, sty, fill=None):
    cell.font, cell.border, cell.alignment, cell.fill = (copy.copy(s) for s in sty[:4])
    cell.number_format = sty[4]
    if fill:
        cell.fill = copy.copy(fill)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--template', required=True)
    ap.add_argument('--sheet', required=True)
    ap.add_argument('--before', required=True, help='Text ô cột A của dòng nhóm sẽ bị đẩy xuống dưới phần chèn')
    ap.add_argument('--data', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--date', required=True, help='Ngày bổ sung DD/MM/YYYY ghi vào Ghi chú')
    args = ap.parse_args()

    rows, techniques = load(args.data)
    errs = validate(rows)
    if errs:
        sys.exit('[LỖI]\n' + '\n'.join(errs))

    wb = openpyxl.load_workbook(args.template)
    ws = wb[args.sheet]
    at = next((r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value or '').strip() == args.before), None)
    if not at:
        sys.exit(f'[LỖI] Không tìm thấy dòng "{args.before}" ở cột A')
    pre_row = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value or '').startswith('Pre-condition'))
    cols = range(1, ws.max_column + 1)
    sty = {'S': {c: style_of(ws.cell(at, c)) for c in cols},
           'P': {c: style_of(ws.cell(pre_row, c)) for c in cols},
           'T': {c: style_of(ws.cell(at - 1, c)) for c in cols}}
    span = {'S': 6, 'P': 7}
    n = len(rows)

    # --- Dời phần dưới điểm chèn xuống n dòng (openpyxl không tự dời merge/công thức/validation/chiều cao) ---
    moved_merges = [CellRange(str(m)) for m in ws.merged_cells.ranges if m.min_row >= at]
    for m in moved_merges:
        ws.unmerge_cells(m.coord)
    heights = {r: ws.row_dimensions[r].height for r in range(at, ws.max_row + 1)}
    old_max = ws.max_row
    ws.insert_rows(at, n)
    for r in range(at + n, old_max + n + 1):
        for c in cols:
            cell = ws.cell(r, c)
            if isinstance(cell.value, str) and cell.value.startswith('='):
                origin = f'{cell.column_letter}{r - n}'
                cell.value = Translator(cell.value, origin=origin).translate_formula(cell.coordinate)
    for r, h in heights.items():
        ws.row_dimensions[r + n].height = h
    for r in range(at, at + n):
        ws.row_dimensions[r].height = None
    for m in moved_merges:
        m.shift(row_shift=n)
        ws.merge_cells(m.coord)

    for dv in ws.data_validations.dataValidation:
        new_ranges = []
        for rng in dv.sqref.ranges:
            cr = CellRange(rng.coord)
            if cr.min_row >= at:
                cr.shift(row_shift=n)
            new_ranges.append(cr.coord)
        dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(' '.join(new_ranges))
    for cf_range in list(ws.conditional_formatting._cf_rules):
        if any(CellRange(r.coord).min_row >= at for r in cf_range.sqref.ranges):
            print('[CẢNH BÁO] Có conditional formatting dưới điểm chèn, chưa dời:', cf_range.sqref)

    # --- Ghi dòng bổ sung ---
    r, group_start, first_tc, keys, tc_rows = at, None, None, {}, []

    def close_group(end):
        nonlocal group_start
        if group_start and end > group_start:
            ws.merge_cells(start_row=group_start, start_column=COL_FUNC, end_row=end, end_column=COL_FUNC)
        group_start = None

    for item in rows:
        kind = item[0]
        for c in cols:
            apply(ws.cell(r, c), sty[kind][c], NEW_FILL if c <= LAST_COL else None)
        if kind in span:
            close_group(r - 1)
            text = indent(item[1])
            ws.cell(r, 1, value=text)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span[kind])
            ws.row_dimensions[r].height = 16 * (text.count('\n') + 1) + 4 if kind == 'P' else None
        else:
            _, func, purpose, steps, expected, note = item[:6]
            first_tc = first_tc or r
            tc_rows.append(r)
            if len(item) == 7 and item[6]:
                keys[item[6]] = r
            ws.cell(r, COL_ID, value=f'=$C$4&"-"&TEXT(counta($E${FIRST_COUNT_ROW}:E{r}),"00")')
            if func:
                close_group(r - 1)
                ws.cell(r, COL_FUNC, value=func)
                group_start = r
            ws.cell(r, COL_PURPOSE, value=purpose)
            ws.cell(r, COL_STEPS, value=indent(steps))
            ws.cell(r, COL_EXP, value=indent(expected))
            full_note = f'[Bổ sung {args.date}] ' + note if note else f'[Bổ sung {args.date}]'
            ws.cell(r, COL_NOTE, value=full_note)
            lines = max(steps.count('\n') + 1 + len(steps) // 45,
                        expected.count('\n') + 1 + len(expected) // 55,
                        len(full_note) // 40 + 1, len(purpose) // 38 + 1)
            ws.row_dimensions[r].height = max(32, min(16 * lines, 300))
        r += 1
    close_group(r - 1)

    dv = DataValidation(type='list', formula1='"Pass,Fail,Pending,N/A"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'F{first_tc}:F{r - 1}')

    # --- ID thật (TC-xx) để sheet kỹ thuật tham chiếu ---
    def tc_id(row):
        cnt = sum(1 for x in range(FIRST_COUNT_ROW, row + 1) if ws.cell(x, COL_EXP).value not in (None, ''))
        return f'TC-{cnt:02d}'

    ids = {k: tc_id(v) for k, v in keys.items()}
    for block in techniques:
        for line in block['rows']:
            for i, v in enumerate(line):
                line[i] = re.sub(r'\{(\w+)\}', lambda mm: ids[mm.group(1)], str(v))

    for name in list(wb.sheetnames):
        if name != ws.title:
            del wb[name]
    if techniques:
        build_technique_sheet(wb, techniques)
    wb.active = 0
    wb.save(args.out)
    print(f'[OK] {args.out}')
    print(f'  Chèn {n} dòng tại dòng {at}..{at + n - 1}: {len(tc_rows)} TC ({tc_id(tc_rows[0])} → {tc_id(tc_rows[-1])})')
    print(f'  Nhóm "{args.before}" dời xuống dòng {at + n}; TC cũ phía sau tăng ID thêm {len(tc_rows)}')


if __name__ == '__main__':
    main()
